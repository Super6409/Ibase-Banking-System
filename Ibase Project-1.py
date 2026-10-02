import openpyxl 
import os

FILE_NAME = "banking_system.xlsx"

# function to create the Excel file if it didint exist already
def setup_excel():
    if not os.path.exists(FILE_NAME):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Accounts"

        sheet["A1"] = "Account Number"
        sheet["B1"] = "Name"
        sheet["C1"] = "Mobile Number"
        sheet["D1"] = "PIN"
        sheet["E1"] = "Balance"

        workbook.save(FILE_NAME)

# function to find account by account number
def find_account(account_number):
    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    for row in range(2, sheet.max_row + 1):
        if str(sheet.cell(row, 1).value) == str(account_number):
            return row

    return None

# function to verify PIN
def verify_pin(row, pin):
    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    # Format cell value as a string (padded to 4 digits if read as an int)
    cell_val = sheet.cell(row, 4).value
    stored_pin = str(cell_val).zfill(4) if cell_val is not None else ""

    if stored_pin == str(pin):
        return True

    return False

# function to create a new account
def create_account():
    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    print("\n--- Create Account ---")

    name = input("Enter your name: ")
    mobile = input("Enter mobile number: ")
    age = input("Enter age: ")
    balance = input("Enter starting balance: ")
    pin = input("Enter 4 digit PIN: ")

    # Explicit string check for PIN (Fix #4)
    pin = str(pin).strip()
    if len(pin) != 4 or not pin.isdigit():
        print("Invalid PIN. PIN must be 4 digits.")
        workbook.close()
        return

    try:
        balance = float(balance)
        age = int(age)
    except:
        print("Invalid input.")
        workbook.close()
        return

    if balance < 0:
        print("Balance cannot be negative.")
        workbook.close()
        return

    # Fix #1: Safely derive the next account number as an integer
    if sheet.max_row == 1:
        account_number = 1001
    else:
        last_val = sheet.cell(sheet.max_row, 1).value
        try:
            account_number = int(last_val) + 1
        except (ValueError, TypeError):
            account_number = 1001

    # Store PIN explicitly as string (Fix #4)
    sheet.append([int(account_number), name, mobile, str(pin), balance])

    workbook.save(FILE_NAME)
    workbook.close()

    print("\nAccount successfully created!")
    print("Your account number is:", account_number)

# function to deposit money
def deposit_money():
    print("\n--- Deposit Money ---")

    account_number = input("Enter account number: ")
    row = find_account(account_number)

    if row is None:
        print("Account not found.")
        return

    pin = input("Enter PIN: ")

    if not verify_pin(row, pin):
        print("Incorrect PIN.")
        return

    amount = input("Enter money to deposit: ")

    try:
        amount = float(amount)
    except:
        print("Invalid amount.")
        return

    if amount <= 0:
        print("Amount must be greater than 0.")
        return

    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    balance = float(sheet.cell(row, 5).value)
    balance = balance + amount

    sheet.cell(row, 5).value = balance

    workbook.save(FILE_NAME)
    workbook.close()

    print(amount, "is successfully deposited.")
    print("Updated balance:", balance)

# function to withdraw money
def withdraw_money():
    print("\n--- Withdraw Money ---")

    account_number = input("Enter account number: ")
    row = find_account(account_number)

    if row is None:
        print("Account not found.")
        return

    pin = input("Enter PIN: ")

    if not verify_pin(row, pin):
        print("Incorrect PIN.")
        return

    amount = input("Enter money to withdraw: ")

    try:
        amount = float(amount)
    except:
        print("Invalid amount.")
        return

    if amount <= 0:
        print("Amount must be greater than 0.")
        return

    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    balance = float(sheet.cell(row, 5).value)

    if amount > balance:
        print("You don't have sufficient balance.")
        workbook.close()
        return

    balance = balance - amount
    sheet.cell(row, 5).value = balance

    workbook.save(FILE_NAME)
    workbook.close()

    print(amount, "is successfully withdrawn.")
    print("Updated balance:", balance)

#function to check account balance
def check_balance():
    print("\n--- Check Balance ---")

    account_number = input("Enter account number: ")
    row = find_account(account_number)

    if row is None:
        print("Account not found.")
        return

    pin = input("Enter PIN: ")

    if not verify_pin(row, pin):
        print("Incorrect PIN.")
        return

    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    name = sheet.cell(row, 2).value
    balance = sheet.cell(row, 5).value

    print("\nAccount Number:", account_number)
    print("Account Holder:", name)
    print("Current Balance:", balance)

    workbook.close()

#main function to run the banking system
def main():
    setup_excel()

    while True:
        print("\n---------------------------------")
        print("       SIMPLE BANKING SYSTEM    ")
        print("---------------------------------")
        print("1. Create Account")
        print("2. Deposit Money")
        print("3. Withdraw Money")
        print("4. Check Balance")
        print("5. Exit")

        choice = input("Enter your choice: ")

        if choice == "1":
            create_account()

        elif choice == "2":
            deposit_money()

        elif choice == "3":
            withdraw_money()

        elif choice == "4":
            check_balance()

        elif choice == "5":
            print("\nThank you for using the banking system!")
            break

        else:
            print("Invalid choice. Please try again.")


main()